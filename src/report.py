"""
report.py
---------
Builds the client-facing Excel report from the cleaned DataFrame.

Sheets:
  Report          - key numbers, top 5 products, orders by city, monthly sales, 3 charts
  Cleaned Data    - the cleaned table (Revenue is a live formula: Qty x Price)
  Cleaning Notes  - what was wrong, what was done, and every assumption

All report numbers are Excel formulas (SUMIFS / COUNTIFS) that read from the
'Cleaned Data' sheet, so the report recalculates if the data is edited.
"""

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
NAVY, TEAL, LIGHT, GREY = "1F3A5F", "1B9E8A", "F3F6FA", "5B6675"
THIN = Side(style="thin", color="D5DBE5")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)

DATA = "'Cleaned Data'"
COLS = ["Order ID", "Order Date", "Customer Name", "Phone", "City", "Product",
        "Category", "Qty", "Price", "Payment Method", "Status", "Month", "Revenue"]
# column letters in 'Cleaned Data'
C = {name: get_column_letter(i + 1) for i, name in enumerate(COLS)}


def _f(bold=False, color="000000", size=10, italic=False):
    return Font(name=FONT, bold=bold, color=color, size=size, italic=italic)


def _header(ws, row, col, labels):
    for i, text in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=text)
        c.font = _f(True, "FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX


def _cell(ws, row, col, value, fmt=None, bold=False, fill=None, align=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font = _f(bold)
    c.border = BOX
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    if align:
        c.alignment = Alignment(horizontal=align)
    return c


def _title(ws, row, text, sub=None):
    ws.cell(row=row, column=1, value=text).font = _f(True, NAVY, 12)
    if sub:
        ws.cell(row=row + 1, column=1, value=sub).font = _f(False, GREY, 9, True)


def _write_data_sheet(wb, df):
    ws = wb.create_sheet("Cleaned Data")
    _header(ws, 1, 1, COLS)
    n = len(df)
    for r, (_, row) in enumerate(df.iterrows(), start=2):
        for j, name in enumerate(COLS, start=1):
            if name == "Revenue":
                q, p = f"{C['Qty']}{r}", f"{C['Price']}{r}"
                val = f'=IF(AND(ISNUMBER({q}),ISNUMBER({p})),{q}*{p},"")'
            elif name == "Month":
                val = str(row["Month"])
            else:
                val = row[name]
                if pd.isna(val):
                    val = None          # left blank on purpose (see Cleaning Notes)
                elif name == "Order Date":
                    val = val.to_pydatetime()
                elif name in ("Qty", "Order ID"):
                    val = int(val)
                elif name == "Price":
                    val = float(val)
            c = ws.cell(row=r, column=j, value=val)
            c.font = _f()
            if name == "Order Date":
                c.number_format = "yyyy-mm-dd"
            elif name in ("Price", "Revenue"):
                c.number_format = "#,##0"
    widths = [10, 12, 22, 14, 13, 12, 13, 6, 9, 17, 11, 9, 10]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{n + 1}"
    return n


def _bar(ws, title, cats, vals, anchor, ytitle, w=15.5, h=7.5, color=TEAL):
    ch = BarChart()
    ch.type = "col"
    ch.title = title
    ch.style = 10
    ch.y_axis.title = ytitle
    ch.height, ch.width = h, w
    ch.add_data(vals, titles_from_data=True)
    ch.set_categories(cats)
    ch.legend = None
    ch.series[0].graphicalProperties.solidFill = color
    ch.dataLabels = DataLabelList()
    ch.dataLabels.showVal = True
    ch.dataLabels.showCatName = False
    ch.dataLabels.showSerName = False
    ch.dataLabels.showLegendKey = False
    ch.dataLabels.showPercent = False
    ch.dataLabels.showBubbleSize = False
    ch.y_axis.delete = False
    ch.x_axis.delete = False
    ws.add_chart(ch, anchor)


def _write_report_sheet(wb, df, n):
    ws = wb.active
    ws.title = "Report"
    ws.sheet_view.showGridLines = False
    rng = lambda col: f"{DATA}!${C[col]}$2:${C[col]}${n + 1}"

    ws["A1"] = "Sales Report - Nour Fashion (Jun - Aug 2026)"
    ws["A1"].font = _f(True, NAVY, 15)
    ws["A2"] = ("Built from the cleaned sales file. Every number below is a formula "
                "that reads the 'Cleaned Data' sheet. Revenue = Qty x Price (gross, "
                "includes cancelled/returned orders).")
    ws["A2"].font = _f(False, GREY, 9, True)

    # ---- Key numbers ----
    _title(ws, 4, "Key numbers")
    _header(ws, 5, 1, ["Metric", "Value"])
    key = [
        ("Total orders", f"=COUNTA({rng('Order ID')})", "#,##0"),
        ("Total revenue (EGP)", f"=SUM({rng('Revenue')})", "#,##0"),
        ("Cancelled + returned orders",
         f'=COUNTIF({rng("Status")},"Cancelled")+COUNTIF({rng("Status")},"Returned")', "#,##0"),
        ("Cancelled + returned share", "=B8/B6", "0.0%"),
    ]
    for i, (label, f, fmt) in enumerate(key, start=6):
        _cell(ws, i, 1, label, fill=LIGHT)
        _cell(ws, i, 2, f, fmt, bold=True)

    # ---- Top 5 products (by revenue) ----
    top = (df.groupby("Product")["Revenue"].sum().sort_values(ascending=False).head(5).index.tolist())
    _title(ws, 11, "Top 5 products by revenue",
           "Ranked by revenue (not units): units are nearly flat across products, "
           "so price is what separates them.")
    _header(ws, 13, 1, ["Rank", "Product", "Revenue (EGP)", "Units sold", "Orders"])
    for i, prod in enumerate(top):
        r = 14 + i
        _cell(ws, r, 1, i + 1, align="center")
        _cell(ws, r, 2, prod)
        _cell(ws, r, 3, f"=SUMIFS({rng('Revenue')},{rng('Product')},B{r})", "#,##0")
        _cell(ws, r, 4, f"=SUMIFS({rng('Qty')},{rng('Product')},B{r})", "#,##0")
        _cell(ws, r, 5, f"=COUNTIFS({rng('Product')},B{r})", "#,##0")

    # ---- Orders by city ----
    cities = df["City"].value_counts().index.tolist()       # sorted by orders
    _title(ws, 20, "Orders by city",
           "Orders with an unknown city are left out of this table (see Cleaning Notes).")
    _header(ws, 22, 1, ["City", "Orders", "Revenue (EGP)"])
    for i, city in enumerate(cities):
        r = 23 + i
        _cell(ws, r, 1, city)
        _cell(ws, r, 2, f"=COUNTIFS({rng('City')},A{r})", "#,##0")
        _cell(ws, r, 3, f"=SUMIFS({rng('Revenue')},{rng('City')},A{r})", "#,##0")
    c0, c1 = 23, 23 + len(cities) - 1
    tr = c1 + 1
    _cell(ws, tr, 1, "Top city (by orders)", bold=True, fill=LIGHT)
    _cell(ws, tr, 2, f"=INDEX(A{c0}:A{c1},MATCH(MAX(B{c0}:B{c1}),B{c0}:B{c1},0))", bold=True, fill=LIGHT)
    _cell(ws, tr, 3, f"=MAX(B{c0}:B{c1})", "#,##0\" orders\"", bold=True, fill=LIGHT)

    # ---- Monthly sales ----
    months = sorted(df["Month"].astype(str).unique())
    m_title = tr + 3
    _title(ws, m_title, "Monthly sales")
    _header(ws, m_title + 1, 1, ["Month", "Orders", "Revenue (EGP)", "Units sold"])
    m0 = m_title + 2
    for i, m in enumerate(months):
        r = m0 + i
        _cell(ws, r, 1, m)
        _cell(ws, r, 2, f"=COUNTIFS({rng('Month')},A{r})", "#,##0")
        _cell(ws, r, 3, f"=SUMIFS({rng('Revenue')},{rng('Month')},A{r})", "#,##0")
        _cell(ws, r, 4, f"=SUMIFS({rng('Qty')},{rng('Month')},A{r})", "#,##0")
    m1 = m0 + len(months) - 1

    for col, w in zip("ABCDE", [30, 16, 16, 12, 10]):
        ws.column_dimensions[col].width = w

    # ---- Charts (to the right of the tables) ----
    _bar(ws, "Top 5 Products by Revenue (EGP)",
         Reference(ws, min_col=2, min_row=14, max_row=18),
         Reference(ws, min_col=3, min_row=13, max_row=18), "G4", "EGP", color=TEAL)
    _bar(ws, "Orders by City",
         Reference(ws, min_col=1, min_row=c0, max_row=c1),
         Reference(ws, min_col=2, min_row=c0 - 1, max_row=c1), "G20", "Orders", color=NAVY)
    _bar(ws, "Monthly Revenue (EGP)",
         Reference(ws, min_col=1, min_row=m0, max_row=m1),
         Reference(ws, min_col=3, min_row=m0 - 1, max_row=m1), "G36", "EGP", color="E8913A")


def _write_notes_sheet(wb, df, raw, n):
    ws = wb.create_sheet("Cleaning Notes")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 78
    ws["A1"] = "Cleaning Notes"
    ws["A1"].font = _f(True, NAVY, 15)
    ws["A2"] = "What was wrong with the raw file, what was done, and every assumption made."
    ws["A2"].font = _f(False, GREY, 9, True)

    dup = int(raw.duplicated().sum())
    _title(ws, 4, "Row counts (from the raw-file audit)")
    _header(ws, 5, 1, ["Item", "Rows", "Note"])
    rows = [
        ("Rows in raw file", len(raw), "Value from the raw Excel file"),
        ("Exact duplicate rows found", dup, "Duplicates are removed AFTER formatting is standardized"),
        ("Rows after cleaning", f"=COUNTA({DATA}!$A$2:$A${n + 1})", "Live count from the 'Cleaned Data' sheet"),
    ]
    for i, (a, b, c) in enumerate(rows, start=6):
        _cell(ws, i, 1, a, fill=LIGHT); _cell(ws, i, 2, b, "#,##0", bold=True); _cell(ws, i, 3, c)

    _title(ws, 10, "Missing values left after cleaning")
    _header(ws, 11, 1, ["Column", "Blank cells", "Decision"])
    miss = [
        ("City", "E", "Left blank. Unmapped values and the '-' placeholder are not guessed. These orders are excluded from the city table only."),
        ("Qty", "H", "Left blank. A guessed quantity would silently change totals."),
        ("Price", "I", "Left blank. Orders without a price have no revenue and are skipped by the revenue formulas."),
    ]
    for i, (col, letter, note) in enumerate(miss, start=12):
        _cell(ws, i, 1, col, fill=LIGHT)
        _cell(ws, i, 2, f"=COUNTBLANK({DATA}!${letter}$2:${letter}${n + 1})", "#,##0", bold=True)
        _cell(ws, i, 3, note)
        ws.cell(row=i, column=3).alignment = Alignment(wrap_text=True, vertical="top")
    _cell(ws, 15, 1, "Orders without revenue", fill=LIGHT)
    _cell(ws, 15, 2, f'=COUNTIF({DATA}!${C["Revenue"]}$2:${C["Revenue"]}${n + 1},"")', "#,##0", bold=True)
    _cell(ws, 15, 3, "Revenue totals are therefore slightly understated, not wrong.")

    _title(ws, 17, "Issues found and how each was handled")
    _header(ws, 18, 1, ["Column", "", "What was done"])
    fixes = [
        ("Order Date", "Mixed formats (ISO, D/M/Y, D-Mon-Y, dotted, real dates) all converted to one date type. Dotted dates can be read two ways; the reading that falls inside the file's date range was used. Ambiguous cases default to day-first."),
        ("City", "Arabic/English spellings, casing and abbreviations mapped to one English name per city."),
        ("Price", "Text such as 'EGP 950', '950 جنيه' and '1,100' converted to numbers."),
        ("Phone", "Unified to the 11-digit local format; '+20' removed; lost leading zeros restored."),
        ("Product / Payment / Status", "Casing and spelling unified (e.g. Tshirt -> T-Shirt, COD -> Cash On Delivery)."),
        ("Qty", "The word 'two' converted to 2; placeholders left blank."),
        ("Customer Name / Phone blanks", "Filled with a readable placeholder ('Not Found' / 'Not provided') because they are never used in a calculation."),
    ]
    for i, (a, c) in enumerate(fixes, start=19):
        _cell(ws, i, 1, a, fill=LIGHT); _cell(ws, i, 2, "")
        _cell(ws, i, 3, c)
        ws.cell(row=i, column=3).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = 42 if len(c) > 95 else 28

    ws["A27"] = ("Data note: this practice project uses a simulated dataset that mimics a real "
                 "messy export, so the business figures are illustrative.")
    ws["A27"].font = _f(False, GREY, 9, True)


def build_excel_report(df, raw, out_path):
    wb = Workbook()
    n = _write_data_sheet(wb, df)
    _write_report_sheet(wb, df, n)
    _write_notes_sheet(wb, df, raw, n)
    wb.move_sheet("Report", offset=-1)
    wb.move_sheet("Cleaning Notes", offset=-1)
    wb._sheets = [wb["Report"], wb["Cleaned Data"], wb["Cleaning Notes"]]
    wb.active = 0
    wb.save(out_path)
    return out_path
